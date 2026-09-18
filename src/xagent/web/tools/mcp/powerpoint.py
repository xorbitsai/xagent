import io
import json
import logging
import os
from typing import Any
from urllib.parse import quote

import requests
from mcp.server.fastmcp import FastMCP
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as PresentationType

from .utils import setup_proxy_env, url_path_id

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("powerpoint-mcp")

setup_proxy_env()

mcp = FastMCP("powerpoint-mcp")

GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
DEFAULT_TIMEOUT_SECONDS = 30
_BINARY_TIMEOUT_SECONDS = 60

_POWERPOINT_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)

# Microsoft Graph exposes a PowerPoint (.pptx) file only as a driveItem
# content blob -- there is no structured "PowerPoint API" resource (no
# per-slide/per-shape Graph endpoints). Every tool here downloads the whole
# file, edits it in memory with python-pptx, and re-uploads the whole file.
# Kept entirely in memory (io.BytesIO, never written to local disk).
_SIMPLE_UPLOAD_MAX_BYTES = 4_000_000

# python-pptx has no public API for adding a slide at other than the layout
# picked, or for removing one at all -- see _delete_slide's own docstring
# for the latter. Layout 1 ("Title and Content") is the conventional
# default for "just add a slide" the same way Word's own UI defaults a new
# document to the Normal style.
_DEFAULT_SLIDE_LAYOUT_INDEX = 1


class _GraphRequestError(RuntimeError):
    """Graph HTTP failure that retains its status without response parsing."""

    def __init__(self, message: str, *, status_code: int) -> None:
        super().__init__(message)
        self.status_code = status_code


def _success(**payload: Any) -> str:
    return json.dumps({"status": "success", **payload}, ensure_ascii=False)


def _error(message: str, *, details: Any = None) -> str:
    payload: dict[str, Any] = {"status": "error", "message": message}
    if details is not None:
        payload["details"] = details
    return json.dumps(payload, ensure_ascii=False)


def _graph_headers(extra_headers: dict[str, str] | None = None) -> dict[str, str]:
    token = os.environ.get("AUTH_TOKEN")
    if not token:
        raise ValueError("AUTH_TOKEN environment variable is missing")

    headers = {
        "Authorization": f"Bearer {token}",
        "Accept": "application/json",
        "Content-Type": "application/json",
    }
    if extra_headers:
        headers.update(extra_headers)
    return headers


def _graph_request(
    method: str,
    path: str,
    *,
    params: dict[str, Any] | None = None,
    body: dict[str, Any] | None = None,
    data: bytes | None = None,
    extra_headers: dict[str, str] | None = None,
    raw: bool = False,
    timeout: int = DEFAULT_TIMEOUT_SECONDS,
) -> Any:
    response = requests.request(
        method=method,
        url=f"{GRAPH_BASE_URL}{path}",
        headers=_graph_headers(extra_headers),
        params=params,
        json=body,
        data=data,
        timeout=timeout,
    )
    try:
        response.raise_for_status()
    except requests.HTTPError as exc:
        response_text = response.text.strip()
        message = str(exc)
        if response_text:
            message = f"{message} - {response_text}"
        raise _GraphRequestError(message, status_code=response.status_code) from exc

    if raw:
        return response.content
    if response.status_code == 204 or not response.content:
        return {}
    return response.json()


def _site_segment(site_id: str) -> str:
    """Percent-encode a caller-supplied Graph site identifier for
    interpolation into a URL path segment.

    A Graph site id is one of: the literal "root", a composite id
    ("hostname,spSiteId,spWebId"), or a "hostname:/server-relative-path"
    form. ':' and '/' stay unescaped because they're structural to the
    third shape, while a '.'/'..' segment is rejected outright -- standard
    HTTP client URL normalization could otherwise walk the request off
    "/sites/{id}/..." and onto a different Graph endpoint under the same
    OAuth token.
    """
    if not isinstance(site_id, str) or not site_id.strip():
        raise ValueError("site_id is required")
    value = site_id.strip()
    if any(segment in (".", "..") for segment in value.split("/")):
        raise ValueError(f"site_id must not contain '.' or '..' segments: {site_id!r}")
    return quote(value, safe=":/,")


def _normalize_relative_path(path: str) -> str:
    """Normalize a drive-relative file path for a root:/{path}: request URL,
    rejecting '.'/'..' segments, a trailing folder separator, and a filename
    ending in a period (Graph/SharePoint's backing storage can silently
    normalize a trailing dot away, so "Deck.pptx." could silently resolve
    to a real, different "Deck.pptx")."""
    value = path.strip().strip("/")
    if not value:
        raise ValueError("file_path is required")
    if path.strip().endswith("/"):
        raise ValueError(
            "file_path must include a filename, not end with a folder separator"
        )
    if "\\" in value:
        raise ValueError("file_path must use '/' separators and must not contain '\\'")
    if any(segment in (".", "..") for segment in value.split("/")):
        raise ValueError(f"file_path must not contain '.' or '..' segments: {path!r}")
    if value.rsplit("/", 1)[-1].endswith("."):
        raise ValueError(f"file_path filename must not end with a period: {path!r}")
    return value


def _item_path(file_path: str, site_id: str | None, drive_id: str | None) -> str:
    normalized = _normalize_relative_path(file_path)
    if site_id:
        site_segment = _site_segment(site_id)
        drive_base = (
            f"/sites/{site_segment}/drives/{url_path_id(drive_id, 'drive_id')}"
            if drive_id
            else f"/sites/{site_segment}/drive"
        )
    elif drive_id:
        drive_base = f"/drives/{url_path_id(drive_id, 'drive_id')}"
    else:
        drive_base = "/me/drive"
    return f"{drive_base}/root:/{quote(normalized, safe='/')}:"


def _content_path(file_path: str, site_id: str | None, drive_id: str | None) -> str:
    return f"{_item_path(file_path, site_id, drive_id)}/content"


def _download_presentation(
    file_path: str, site_id: str | None, drive_id: str | None
) -> PresentationType:
    content = _graph_request(
        "GET", _content_path(file_path, site_id, drive_id), raw=True
    )
    try:
        return Presentation(io.BytesIO(content))
    except Exception as exc:
        raise ValueError(
            f"Could not open {file_path!r} as a PowerPoint presentation -- it may "
            "not be a valid .pptx file"
        ) from exc


def _upload_presentation(
    presentation: PresentationType,
    file_path: str,
    site_id: str | None,
    drive_id: str | None,
) -> dict[str, Any]:
    buffer = io.BytesIO()
    presentation.save(buffer)
    content = buffer.getvalue()
    if len(content) > _SIMPLE_UPLOAD_MAX_BYTES:
        raise ValueError(
            f"The updated presentation is {len(content)} bytes, over the "
            f"{_SIMPLE_UPLOAD_MAX_BYTES // 1_000_000} MB limit this tool currently "
            "supports"
        )
    result = _graph_request(
        "PUT",
        _content_path(file_path, site_id, drive_id),
        extra_headers={"Content-Type": _POWERPOINT_MIME_TYPE},
        data=content,
        timeout=_BINARY_TIMEOUT_SECONDS,
    )
    if not isinstance(result, dict) or not result.get("id"):
        raise RuntimeError("Graph did not confirm the presentation upload completed")
    safe_item = dict(result)
    safe_item.pop("@microsoft.graph.downloadUrl", None)
    return safe_item


def _create_only_upload(
    content: bytes, file_path: str, site_id: str | None, drive_id: str | None
) -> dict[str, Any]:
    """Upload content, failing instead of silently overwriting if a file
    already exists at file_path.

    A separate exists-check GET followed by a plain _upload_presentation PUT
    would still race: Graph's own "Upload small files" reference documents
    no If-None-Match or other conditional-write option on the simple
    content PUT, so two concurrent callers can both pass the check and the
    second would silently clobber the first. The upload-session API is
    documented to accept a conflictBehavior of "fail" on session creation,
    which Graph enforces atomically server-side (a 409 nameAlreadyExists if
    the target already exists) -- used here even though the content (a
    blank new presentation) is always small enough for a single-PUT
    session, specifically for that atomicity guarantee.
    """
    item_path = _item_path(file_path, site_id, drive_id)
    try:
        session = _graph_request(
            "POST",
            f"{item_path}/createUploadSession",
            body={"item": {"@microsoft.graph.conflictBehavior": "fail"}},
        )
    except _GraphRequestError as exc:
        if exc.status_code == 409:
            raise ValueError(
                f"{file_path!r} already exists; use the other powerpoint_* "
                "tools to edit it instead of recreating it"
            ) from exc
        raise
    upload_url = session.get("uploadUrl") if isinstance(session, dict) else None
    if not isinstance(upload_url, str) or not upload_url:
        raise RuntimeError("Graph did not return an upload session URL")

    # The upload-session URL is itself pre-authenticated (a token in its own
    # query string) -- Graph's createUploadSession docs warn that including
    # an Authorization header on this PUT can cause a 401 -- so this goes
    # through a plain requests.put, not _graph_request (which always
    # attaches one). Both requests.put's own exception (a connection failure
    # or timeout) and the HTTPError from raise_for_status() embed the full
    # request URL in their default str(), so neither is stringified into a
    # message below, and each is re-raised with "from None" rather than
    # "from exc" -- chaining the original would still attach it as
    # __cause__, which a future traceback/log/APM capture could surface --
    # matching onedrive.py's identical guard on the same hazard.
    try:
        response = requests.put(
            upload_url,
            data=content,
            headers={
                "Content-Range": f"bytes 0-{len(content) - 1}/{len(content)}",
                "Content-Type": _POWERPOINT_MIME_TYPE,
            },
            timeout=_BINARY_TIMEOUT_SECONDS,
        )
        response.raise_for_status()
    except requests.HTTPError:
        status_code = response.status_code
        if status_code == 409:
            raise ValueError(
                f"{file_path!r} already exists; use the other powerpoint_* "
                "tools to edit it instead of recreating it"
            ) from None
        raise _GraphRequestError(
            f"PowerPoint presentation upload failed with HTTP {status_code}",
            status_code=status_code,
        ) from None
    except requests.RequestException:
        raise RuntimeError("PowerPoint presentation upload failed") from None
    result = response.json()
    if not isinstance(result, dict) or not result.get("id"):
        raise RuntimeError("Graph did not confirm the presentation upload completed")
    safe_item = dict(result)
    safe_item.pop("@microsoft.graph.downloadUrl", None)
    return safe_item


def _shape_text(shape: Any) -> str | None:
    return shape.text_frame.text if shape.has_text_frame else None


def _select_body_placeholder(slide: Any) -> Any | None:
    """Find the first non-title placeholder on slide that can hold text.

    A placeholder's idx alone doesn't guarantee it has a text frame -- a
    picture, chart, or other content placeholder can sit at a lower idx
    than a real text placeholder (verified: python-pptx's own
    SlidePlaceholders iterates in idx order, with no guarantee idx order
    matches "text placeholders first"), so this skips non-text
    placeholders entirely rather than stopping at the first non-title one
    regardless of whether it can actually hold text.
    """
    return next(
        (
            ph
            for ph in slide.placeholders
            if ph.placeholder_format.idx != 0 and ph.has_text_frame
        ),
        None,
    )


def _require_int(value: Any, field_name: str) -> int:
    """Validate an integer-typed tool argument.

    FastMCP/Pydantic validates and coerces arguments against a tool's type
    hints for a real call over the MCP protocol, but that validation is
    bypassed by a caller that invokes this Python function directly (e.g. a
    test, or any other in-process caller) -- without this, a non-int index
    reaches python-pptx's own indexing/comparison and fails with a raw,
    less clear TypeError instead. bool is excluded even though it's a
    subclass of int in Python, since True/False are never valid indices.
    """
    if not isinstance(value, int) or isinstance(value, bool):
        raise TypeError(f"{field_name} must be an integer")
    return value


def _text_frame_has_dynamic_content(text_frame: Any) -> bool:
    """Whether text_frame holds a hyperlink or a dynamic field (e.g. an
    auto-updating slide number or date placeholder) that a full-frame text
    replacement would silently discard.

    python-pptx's TextFrame.text setter clears every existing paragraph
    and rebuilds the frame from a single new run per line (verified
    directly: assigning .text removes all <a:p> children) -- a run's
    <a:hlinkClick> and a paragraph-level <a:fld> element are both lost
    with no error or warning.
    """
    tx_body = text_frame._txBody
    return (
        tx_body.find(".//" + qn("a:hlinkClick")) is not None
        or tx_body.find(".//" + qn("a:fld")) is not None
    )


def _require_slide(presentation: PresentationType, slide_index: int) -> Any:
    slides = presentation.slides
    if not 0 <= slide_index < len(slides):
        raise ValueError(
            f"slide_index {slide_index} is out of range for a presentation with "
            f"{len(slides)} slides"
        )
    return slides[slide_index]


def _delete_slide(presentation: PresentationType, slide_index: int) -> None:
    """Remove a slide by index.

    python-pptx has no public API for removing a slide (confirmed: no
    Slides.remove()/delete() method exists in the library). The documented
    community workaround -- dropping the slide's <p:sldId> entry from the
    presentation's slide-id list, which is what actually determines slide
    order/membership in the underlying XML -- is used here instead.

    Removing only the <p:sldId> entry leaves the slide's own XML part and
    the presentation-to-slide relationship in the saved package -- verified
    directly: python-pptx's serializer does no orphan-pruning on save, so a
    caller reading the package's parts/relationships directly (rather than
    walking presentation.slides) would still see the "deleted" slide's
    content, and repeated add/delete cycles would accumulate dead parts.
    part.drop_rel() below removes that relationship (and, since it's the
    part's only remaining reference, the part itself) so the deleted
    slide's content doesn't survive in the saved file at all.
    """
    slide_id_list = presentation.slides._sldIdLst
    slide_ids = list(slide_id_list)
    if not 0 <= slide_index < len(slide_ids):
        raise ValueError(
            f"slide_index {slide_index} is out of range for a presentation with "
            f"{len(slide_ids)} slides"
        )
    target = slide_ids[slide_index]
    relationship_id = target.get(qn("r:id"))
    slide_id_list.remove(target)
    if relationship_id:
        presentation.part.drop_rel(relationship_id)


@mcp.tool()
def powerpoint_create_presentation(
    file_path: str, site_id: str | None = None, drive_id: str | None = None
) -> str:
    """Create a new, blank PowerPoint presentation at file_path. Fails if a
    file already exists there -- edit it with the other powerpoint_* tools
    instead of recreating it."""
    try:
        buffer = io.BytesIO()
        Presentation().save(buffer)
        item = _create_only_upload(buffer.getvalue(), file_path, site_id, drive_id)
        return _success(item=item)
    except Exception as e:
        logger.error("Error creating PowerPoint presentation %s: %s", file_path, e)
        return _error(str(e))


@mcp.tool()
def powerpoint_get_presentation_text(
    file_path: str, site_id: str | None = None, drive_id: str | None = None
) -> str:
    """Get all text in a PowerPoint presentation, as a list of slides each
    with the text of every text-bearing shape on it."""
    try:
        presentation = _download_presentation(file_path, site_id, drive_id)
        slides = [
            {
                "slide_index": slide_index,
                "shapes": [
                    text
                    for shape in slide.shapes
                    if (text := _shape_text(shape)) is not None and text
                ],
            }
            for slide_index, slide in enumerate(presentation.slides)
        ]
        return _success(slides=slides)
    except Exception as e:
        logger.error(
            "Error getting text for PowerPoint presentation %s: %s", file_path, e
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_list_slides(
    file_path: str, site_id: str | None = None, drive_id: str | None = None
) -> str:
    """List a PowerPoint presentation's slides, with each slide's index,
    layout name, and shape count."""
    try:
        presentation = _download_presentation(file_path, site_id, drive_id)
        slides = [
            {
                "slide_index": slide_index,
                "layout_name": slide.slide_layout.name,
                "shape_count": len(slide.shapes),
            }
            for slide_index, slide in enumerate(presentation.slides)
        ]
        return _success(slides=slides)
    except Exception as e:
        logger.error(
            "Error listing slides for PowerPoint presentation %s: %s", file_path, e
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_get_slide_text(
    file_path: str,
    slide_index: int,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Get one slide's shapes with their index, type, and text (needed by
    powerpoint_set_shape_text)."""
    try:
        slide_index = _require_int(slide_index, "slide_index")
        presentation = _download_presentation(file_path, site_id, drive_id)
        slide = _require_slide(presentation, slide_index)
        shapes = [
            {
                "shape_index": shape_index,
                "shape_type": str(shape.shape_type),
                "is_placeholder": shape.is_placeholder,
                "text": _shape_text(shape),
            }
            for shape_index, shape in enumerate(slide.shapes)
        ]
        return _success(shapes=shapes)
    except Exception as e:
        logger.error(
            "Error getting slide %s text for PowerPoint presentation %s: %s",
            slide_index,
            file_path,
            e,
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_set_shape_text(
    file_path: str,
    slide_index: int,
    shape_index: int,
    text: str,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Replace a shape's text on a slide, by slide_index and shape_index
    (from powerpoint_get_slide_text). Only works on a shape that already
    has a text frame (a title, body, or plain text box); errors otherwise.
    Also errors (rather than silently discarding it) if the shape's text
    contains a hyperlink or a dynamic field such as an auto-updating slide
    number or date, since replacing the whole frame's text has no way to
    carry those over -- edit that shape directly in PowerPoint instead."""
    try:
        slide_index = _require_int(slide_index, "slide_index")
        shape_index = _require_int(shape_index, "shape_index")
        presentation = _download_presentation(file_path, site_id, drive_id)
        slide = _require_slide(presentation, slide_index)
        shapes = list(slide.shapes)
        if not 0 <= shape_index < len(shapes):
            raise ValueError(
                f"shape_index {shape_index} is out of range for slide {slide_index} "
                f"with {len(shapes)} shapes"
            )
        shape = shapes[shape_index]
        if not shape.has_text_frame:
            raise ValueError(
                f"shape {shape_index} on slide {slide_index} has no text frame "
                "(e.g. it's an image or a plain line/connector) and cannot hold text"
            )
        if _text_frame_has_dynamic_content(shape.text_frame):
            raise ValueError(
                f"shape {shape_index} on slide {slide_index} contains a hyperlink "
                "or a dynamic field (e.g. slide number or date) that replacing its "
                "text would silently delete -- edit this shape directly in "
                "PowerPoint instead"
            )
        shape.text_frame.text = text
        item = _upload_presentation(presentation, file_path, site_id, drive_id)
        return _success(item=item)
    except Exception as e:
        logger.error(
            "Error setting shape %s text on slide %s in PowerPoint presentation %s: %s",
            shape_index,
            slide_index,
            file_path,
            e,
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_add_slide(
    file_path: str,
    title: str | None = None,
    body_text: str | None = None,
    layout_index: int = _DEFAULT_SLIDE_LAYOUT_INDEX,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Add a new slide at the end of a PowerPoint presentation.

    layout_index selects a slide layout from the presentation's slide
    master (0 is usually "Title Slide", 1 "Title and Content", 6 "Blank" --
    use powerpoint_list_slide_layouts to see what this presentation
    actually has, since layouts vary by template). title, if given, is set
    on the layout's title placeholder when present. body_text, if given, is
    set on the first non-title placeholder (by idx order) that has a text
    frame -- a non-text placeholder earlier in idx order (e.g. a picture or
    chart placeholder) is skipped rather than causing body_text to be left
    unset. This is still a heuristic, not a guaranteed "body" role: on a
    "Title Slide" layout, for example, that placeholder is actually the
    subtitle. Use powerpoint_get_slide_text after adding the slide to
    confirm where each piece of text actually landed."""
    try:
        layout_index = _require_int(layout_index, "layout_index")
        presentation = _download_presentation(file_path, site_id, drive_id)
        layouts = presentation.slide_layouts
        if not 0 <= layout_index < len(layouts):
            raise ValueError(
                f"layout_index {layout_index} is out of range for this "
                f"presentation's {len(layouts)} slide layouts"
            )
        slide = presentation.slides.add_slide(layouts[layout_index])
        if title is not None and slide.shapes.title is not None:
            slide.shapes.title.text = title
        if body_text is not None:
            body_placeholder = _select_body_placeholder(slide)
            if body_placeholder is not None:
                body_placeholder.text_frame.text = body_text
        item = _upload_presentation(presentation, file_path, site_id, drive_id)
        return _success(item=item, slide_index=len(presentation.slides) - 1)
    except Exception as e:
        logger.error(
            "Error adding slide to PowerPoint presentation %s: %s", file_path, e
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_list_slide_layouts(
    file_path: str, site_id: str | None = None, drive_id: str | None = None
) -> str:
    """List the slide layouts available in a presentation's template, with
    the index powerpoint_add_slide's layout_index expects."""
    try:
        presentation = _download_presentation(file_path, site_id, drive_id)
        layouts = [
            {"layout_index": index, "name": layout.name}
            for index, layout in enumerate(presentation.slide_layouts)
        ]
        return _success(layouts=layouts)
    except Exception as e:
        logger.error(
            "Error listing slide layouts for PowerPoint presentation %s: %s",
            file_path,
            e,
        )
        return _error(str(e))


@mcp.tool()
def powerpoint_delete_slide(
    file_path: str,
    slide_index: int,
    site_id: str | None = None,
    drive_id: str | None = None,
) -> str:
    """Delete a slide from a PowerPoint presentation by index."""
    try:
        slide_index = _require_int(slide_index, "slide_index")
        presentation = _download_presentation(file_path, site_id, drive_id)
        _delete_slide(presentation, slide_index)
        item = _upload_presentation(presentation, file_path, site_id, drive_id)
        return _success(item=item)
    except Exception as e:
        logger.error(
            "Error deleting slide %s from PowerPoint presentation %s: %s",
            slide_index,
            file_path,
            e,
        )
        return _error(str(e))


if __name__ == "__main__":
    mcp.run()
